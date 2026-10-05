import os
import base64
import logging
from typing import List, Dict, Any

from pptx import Presentation
from openai import AzureOpenAI

logger = logging.getLogger(__name__)


# ── Slide-level extractors ────────────────────────────────────────

def extract_slide_text(slide) -> str:
    """Extract all text from a slide's shapes (excluding tables)."""
    texts = []
    for shape in slide.shapes:
        if shape.has_table:
            continue  # handled separately
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                text = paragraph.text.strip()
                if text:
                    texts.append(text)
    return "\n".join(texts)


def extract_slide_tables(slide) -> str:
    """Extract table content from a slide as pipe-delimited rows."""
    tables_text = []
    for shape in slide.shapes:
        if shape.has_table:
            table = shape.table
            rows = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                rows.append(" | ".join(cells))
            tables_text.append("\n".join(rows))
    return "\n\n".join(tables_text)


def extract_slide_images(slide) -> List[tuple]:
    """Extract images from a slide as (bytes, content_type) tuples."""
    images = []
    for shape in slide.shapes:
        if hasattr(shape, "image"):
            try:
                blob = shape.image.blob
                content_type = shape.image.content_type or "image/png"
                images.append((blob, content_type))
            except Exception:
                pass
    return images


def describe_image(
    client: AzureOpenAI, deployment: str, image_bytes: bytes, content_type: str
) -> str:
    """Send an image to GPT-4o vision and return a text description."""
    b64 = base64.b64encode(image_bytes).decode("utf-8")
    data_uri = f"data:{content_type};base64,{b64}"

    try:
        response = client.chat.completions.create(
            model=deployment,
            messages=[
                {
                    "role": "system",
                    "content": (
                        "Describe this image from a presentation slide concisely. "
                        "Focus on key information, data, labels, and any text visible."
                    ),
                },
                {
                    "role": "user",
                    "content": [
                        {
                            "type": "image_url",
                            "image_url": {"url": data_uri, "detail": "high"},
                        }
                    ],
                },
            ],
            max_tokens=500,
        )
        return response.choices[0].message.content
    except Exception as e:
        logger.warning(f"Failed to describe image: {e}")
        return "[Image could not be described]"


# ── PPT-level processing ──────────────────────────────────────────

def process_ppt(
    file_path: str, client: AzureOpenAI, deployment: str
) -> Dict[str, Any]:
    """Process a single PPTX file and return structured slide content."""
    filename = os.path.basename(file_path)
    prs = Presentation(file_path)
    slides_content = []

    for idx, slide in enumerate(prs.slides, start=1):
        text = extract_slide_text(slide)
        tables = extract_slide_tables(slide)
        images = extract_slide_images(slide)

        # Describe each image via GPT-4o vision
        image_descriptions = []
        for img_bytes, ctype in images:
            desc = describe_image(client, deployment, img_bytes, ctype)
            image_descriptions.append(desc)

        # Combine all content for this slide
        parts = []
        if text:
            parts.append(f"Text:\n{text}")
        if tables:
            parts.append(f"Tables:\n{tables}")
        if image_descriptions:
            parts.append("Images:\n" + "\n".join(image_descriptions))

        combined = "\n\n".join(parts) if parts else "[Empty slide]"

        slides_content.append(
            {"slide_number": idx, "content": combined}
        )

    logger.info(f"Processed {filename}: {len(slides_content)} slides")
    return {
        "filename": filename,
        "file_path": file_path,
        "total_slides": len(slides_content),
        "slides": slides_content,
    }


def process_all_ppts(
    volume_path: str, client: AzureOpenAI, deployment: str
) -> List[Dict[str, Any]]:
    """Process every .pptx file in the given volume directory."""
    ppt_data = []
    for fname in sorted(os.listdir(volume_path)):
        if fname.lower().endswith(".pptx"):
            fpath = os.path.join(volume_path, fname)
            ppt_data.append(process_ppt(fpath, client, deployment))
    return ppt_data
